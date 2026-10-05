from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from pathlib import Path

out=Path(r'F:\Versha Ekrocx\live-projects\ecomm\docs\Ecommerce Initial Requirements.docx')
doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.62);sec.left_margin=sec.right_margin=Inches(.7)
sec.footer_distance=Inches(.25)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','List Bullet','List Number']:
    style=doc.styles[name];style.font.name='Calibri';style.font.color.rgb=RGBColor(0,0,0)
    style.paragraph_format.space_after=Pt(4)
normal=doc.styles['Normal'];normal.font.size=Pt(10);normal.paragraph_format.line_spacing=1.05
for name,size in [('Title',23),('Subtitle',11),('Heading 1',14),('Heading 2',11)]:
    s=doc.styles[name];s.font.size=Pt(size);s.font.bold=name!='Subtitle';s.paragraph_format.space_before=Pt(9 if name.startswith('Heading') else 0);s.paragraph_format.space_after=Pt(5);s.paragraph_format.keep_with_next=True
for name in ['List Bullet','List Number']:
    doc.styles[name].font.size=Pt(10);doc.styles[name].paragraph_format.space_after=Pt(3)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('Initial requirements  |  ');r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('777777')
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
doc.core_properties.title='Ecommerce Project Initial Requirements'
doc.core_properties.subject='Initial technology setup and functional scope'
doc.core_properties.author=''

def p(text,style=None): return doc.add_paragraph(text,style)
def h(text,level=1): doc.add_heading(text,level)
def bullet(text): p(text,'List Bullet')
def entry(label,text,number=None):
    para=doc.add_paragraph();para.paragraph_format.space_after=Pt(4)
    para.add_run((str(number)+'. ' if number else '')+label+' — ').bold=True;para.add_run(text)

p('Ecommerce Project Initial Requirements','Title')
p('Custom ecommerce platform for Kärcher Midea and LG','Subtitle')
p('This document defines the initial technology setup and functional scope for a custom ecommerce platform, comprising a customer storefront, a connected admin panel and a shared backend API.')
h('1 Technology stack')
table=doc.add_table(rows=1,cols=2);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
table.columns[0].width=Inches(1.42);table.columns[1].width=Inches(5.45)
for cell,text in zip(table.rows[0].cells,['Component','Requirement']):
    cell.text=text
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7E7E7');cell._tc.get_or_add_tcPr().append(shade)
    for r in cell.paragraphs[0].runs:r.bold=True
rows=[
('Backend and API','Laravel using PHP 8.x and the latest stable Laravel release.'),
('Storefront','React with Next.js and server rendering for fast product pages and search engine visibility.'),
('Admin panel','React connected to the Laravel API.'),
('Database','MySQL.'),
('Performance','Redis caching, Cloudflare CDN and optimised images.'),
('Payments','M-Pesa through the Daraja API with STK Push, plus a card gateway such as Pesapal, Flutterwave or DPO. The card gateway is to be confirmed.'),
('Email','Transactional email delivery through SMTP or Amazon SES.'),
('Security','SSL, role-based admin access, secure password storage and protection against common web attacks.')]
for a,b in rows:
    cells=table.add_row().cells;cells[0].text=a;cells[1].text=b
for row in table.rows:
    trPr=row._tr.get_or_add_trPr();trPr.append(OxmlElement('w:cantSplit'))
    for cell in row.cells:
        cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        pr=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
        for side in ['top','left','bottom','right']:
            el=OxmlElement('w:'+side);el.set(qn('w:w'),'85');el.set(qn('w:type'),'dxa');m.append(el)
        pr.append(m)
        borders=OxmlElement('w:tcBorders')
        for side in ['top','left','bottom','right']:
            el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
        pr.append(borders)
        for para in cell.paragraphs:para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=1.0
h('2 Custom development approach')
p('The platform will use fully custom design and development rather than WordPress or WooCommerce.')
for text in ['A unique design built around the brands rather than a ready-made theme.','No dependency on third-party ecommerce plugins, reducing plugin security risks and update conflicts.','A clean, lightweight codebase designed for faster performance.','Full ownership of the source code and freedom to add features as the business grows.']:bullet(text)
h('3 Admin panel scope')
h('Dashboard',2)
bullet('Sales, orders, new customers and average order value for today, this week and this month.')
bullet('Recent orders, pending payments and low-stock alerts.')
bullet('Sales trend chart and top-selling products.')

doc.add_page_break()
h('Admin management and reporting')
h('Master modules',2)
masters=[('Brands','Kärcher, Midea, LG and future brands, with logo, banner and description.'),('Categories','Multi-level categories and subcategories.'),('Attributes and specifications','Capacity, power, size, colour and other details for filters and specification tables.'),('Products','Name, SKU, price, sale price, images, specifications, warranty information, PDF manuals or brochures and SEO fields.'),('Stock and inventory','Stock quantity per product and low-stock thresholds.'),('Customers','Profiles, addresses, order history and account enable or disable controls.'),('Delivery zones and rates','Counties or towns, delivery charges and free-delivery rules.'),('Tax settings','VAT configuration for Kenya.'),('Coupons and discounts','Percentage or fixed discounts, validity dates and usage limits.'),('Banners and homepage content','Sliders, promotional banners, featured products and featured brands.'),('CMS pages','About, Contact, FAQ, Shipping and Returns, Privacy and Terms.'),('Admin users and roles','Multiple staff logins with permissions, including Admin, Sales and Store Manager.')]
for i,(label,text) in enumerate(masters,1):entry(label,text,i)
h('Operations modules',2)
for label,text in [('Orders','View orders; update status from Pending to Confirmed, Dispatched and Delivered, or Cancelled; print invoices and add notes.'),('Payments','M-Pesa and card transaction logs, payment statuses and failed payments.'),('Returns and refunds','Record return requests and refund statuses.'),('Enquiries','Manage contact-form and product enquiries.'),('Email notifications','Order confirmations, order status updates and admin new-order alerts.'),('Settings','Store details, logo, contact information, social links, payment settings and email settings.')]:entry(label,text)
h('Reports and analytics',2)
p('All reports must support date-range filtering and export to Excel or CSV.')
for text in ['Sales totals grouped daily, weekly and monthly; sales by brand and category.','Top-selling products; orders by status including pending, delivered and cancelled.','Payments by M-Pesa or card and successful or failed outcome.','New and returning customers, top customers, low-stock and out-of-stock products.','Coupon usage report.']:bullet(text)
p('Google Analytics 4 and Meta Facebook Pixel integration will support traffic, visitor behaviour and marketing tracking.')

doc.add_page_break()
h('4 Front store scope')
sections=[('Header and navigation','Menus by brand and category, a search bar with live suggestions, and cart and account icons.'),('Homepage','Hero banner slider; Kärcher, Midea and LG brand showcase; featured categories; featured, best-selling and new products; promotional banners; newsletter signup.'),('Brand pages','A dedicated page for each brand, including a banner, introduction and all of its products.'),('Category and product listings','Filters by brand, category, price and specifications; sorting by price, newest and popularity; pagination.'),('Product detail page','Image gallery with zoom; regular and offer prices; stock status; key specifications table; description; warranty information; downloadable brochure or manual; related products; WhatsApp enquiry button.'),('Cart and checkout','Add, update and remove items; apply coupons; calculate delivery charges by location; support guest checkout or customer login. Accept M-Pesa STK Push or card payments and provide an order confirmation page and confirmation email.'),('Customer account','Registration and login, order history and status tracking, saved addresses, wishlist, profile management and password management.'),('Information pages','About, Contact with form, map and WhatsApp, FAQ, Shipping and Returns, Privacy Policy, and Terms and Conditions.'),('General requirements','Responsive layouts for mobile, tablet and desktop; SEO-ready clean URLs and meta tags; fast loading through caching, CDN and optimised images; WhatsApp chat button and social media links.')]
for label,text in sections:h(label,2);p(text)
doc.save(out)
print(out)
