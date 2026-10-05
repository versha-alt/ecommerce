from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(r'F:\Versha Ekrocx\live-projects\ecomm')
out=root/'docs/Ecommerce Consolidated Project Requirements.docx'
doc=Document();s=doc.sections[0];s.page_width=Inches(8.27);s.page_height=Inches(11.69)
s.top_margin=s.bottom_margin=Inches(.55);s.left_margin=s.right_margin=Inches(.65);s.footer_distance=Inches(.24)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','List Bullet']:
    style=doc.styles[name];style.font.name='Calibri';style.font.color.rgb=RGBColor(0,0,0);style.font.italic=False
    style.paragraph_format.space_after=Pt(4);style.paragraph_format.line_spacing=1.0
    if name=='Normal':style.font.size=Pt(10)
for name,size in [('Title',21),('Subtitle',10),('Heading 1',13),('Heading 2',11)]:
    style=doc.styles[name];style.font.size=Pt(size);style.font.bold=name!='Subtitle';style.paragraph_format.space_before=Pt(7 if name.startswith('Heading') else 0);style.paragraph_format.space_after=Pt(4);style.paragraph_format.keep_with_next=True
for el in doc.styles._element.xpath('.//w:pBdr'):el.getparent().remove(el)
foot=s.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=foot.add_run('Project requirements  |  ');r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('777777')
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');foot._p.append(f)
doc.core_properties.title='Ecommerce Consolidated Project Requirements';doc.core_properties.subject='Initial system specification';doc.core_properties.author=''
def p(text):return doc.add_paragraph(text)
def h(text):doc.add_heading(text,1)
def e(label,text):
    par=p('');par.add_run(label+' — ').bold=True;par.add_run(text)

doc.add_paragraph('Ecommerce Project Requirements','Title')
doc.add_paragraph('Consolidated initial specification','Subtitle')
p('The system shall provide a custom ecommerce storefront and an admin panel for Kärcher, Midea, LG and future brands. Olive green shall be the primary design colour. The following requirements define the complete technology setup, modules and customer workflows.')
h('1 Technology and development approach')
for label,text in [
('Backend and API','Laravel on PHP 8.x using the latest stable Laravel release; a shared API shall connect the storefront and admin panel.'),
('Storefront and admin','The storefront shall use React with Next.js and server rendering for fast, search-friendly product pages. The admin panel shall use React connected to Laravel.'),
('Database and performance','MySQL shall persist system data. Redis caching, Cloudflare CDN and optimised images shall support fast loading.'),
('Payments','Support M-Pesa through Daraja STK Push and a card gateway such as Pesapal, Flutterwave or DPO, with the card provider to be confirmed. Payment-method configuration shall also support PayPal, Razorpay and other online gateways, plus COD and other manual or offline methods.'),
('Email and security','Use SMTP or Amazon SES for transactional emails. Require SSL, secure password storage, role-based admin access and protection against common web attacks.')]:e(label,text)
p('Design and development shall be fully custom rather than WordPress or WooCommerce. The design shall reflect the brands rather than use a ready-made theme, avoid dependence on third-party ecommerce plugins and their security or update conflicts, and use a clean, lightweight codebase. The business shall own the source code and be able to extend features as it grows.')
h('2 Admin dashboard catalog and configuration')
e('Dashboard','Show sales, orders, new customers and average order value for today, this week and this month; recent orders, pending payments and low-stock alerts; a sales trend chart and top-selling products.')
for label,text in [
('Brands','Manage Kärcher, Midea, LG and future brands, including logos, banners and descriptions.'),
('Categories','Manage multi-level categories and subcategories. Assigning a product to a child shall also associate it with every applicable parent, for example Smartphones → Mobiles → Electronics. The backend and database shall maintain these associations consistently.'),
('Attributes and specifications','Manage reusable details such as capacity, power, size and colour for storefront filters and product specification tables.'),
('Products','Manage name, SKU, regular and sale prices, images, specifications, warranty information, PDF manuals or brochures, and SEO fields. Each product slug shall be unique; backend validation and a database constraint shall prevent duplicates on creation and update.'),
('Product CSV import','Provide CSV file import within Products. Imported products shall follow the same validation, unique-slug and category-association rules as products entered through the normal product workflow.'),
('Stock and inventory','Manage quantities per product and low-stock thresholds.'),
('Delivery and tax','Configure counties or towns, delivery charges and free-delivery rules. Provide VAT configuration for Kenya.'),
('Admin users and roles','Provide multiple staff logins with permissions, including Admin, Sales and Store Manager.')]:e(label,text)

doc.add_page_break()
h('3 Customer commerce and discount workflows')
e('Customers','Customers shall originate from the storefront. The admin panel shall have no Add Customer option. Admins shall manage existing profiles, addresses, order histories and account enable or disable settings without duplicating customer creation.')
e('Orders','Customers shall place orders through storefront checkout; the admin shall have no Add Order option. Admins shall view orders, add internal notes, print invoices and update status through Pending → Confirmed → Dispatched → Delivered, or Cancelled as applicable. Status changes shall persist in the database and appear correctly in the storefront. Maintain an order status history or audit trail where the architecture supports it.')
e('Payment methods','Keep payment configuration separate from transaction history. Group methods into Online Payments, including PayPal, Razorpay, M-Pesa and other required gateways, and Manual or Offline Payments, including COD and configurable alternatives. Admins shall enable or disable methods and manage relevant gateway configuration securely.')
e('Payment history and transactions','Payments shall originate from customer checkout when an order is placed and payment is completed. Transactions shall be created automatically by the payment flow; neither an Add Payment option nor manual payment creation through the admin or backend shall be permitted. Admins shall view and manage existing records. Show Payment ID, Order ID, customer, payment method, amount, payment status, transaction or reference ID and date. Support Pending, Paid or Successful, Failed, Refunded and Partially Refunded where applicable. Payment events, order payment state and refund outcomes shall remain consistent across APIs, the database, storefront and admin panel.')
e('Returns and refunds','Customers shall submit return and refund requests through the storefront UI. The admin module shall have no Add option and shall display and manage existing requests only. Admins shall view request details, review requests, approve or reject them, update relevant statuses, and process or manage refunds where applicable. Refund status shall be reflected in the related payment history.')
e('Coupons and discounts','Use Shopify as a reference for discount creation and management, while adapting the flow to this application rather than copying Shopify’s UI. Provide four distinct discount types:')
for label,text in [
('Amount off products','Percentage or fixed discounts applying to eligible products or categories.'),
('Amount off order','Percentage or fixed discounts on an eligible order.'),
('Buy X get Y','Define the qualifying purchase and the products or quantities rewarded by the offer.'),
('Free shipping','Remove eligible delivery charges under the configured offer conditions.')]:e(label,text)
p('The creation flow shall capture the discount value and value type, customer eligibility, minimum purchase requirements, product or category restrictions, usage limits, start and end dates, and active or inactive status as appropriate to each discount type. The backend shall enforce applicable conditions when an order is calculated.')
e('Integration consistency','Changes shall use the current database models, APIs, frontend and admin architecture. Avoid duplicate functionality and preserve the product, order, customer and settings workflows while enforcing the customer-originated creation rules above.')

doc.add_page_break()
h('4 Content communications and reporting')
for label,text in [
('Homepage content and CMS','Manage sliders, promotional banners, featured products and featured brands. Maintain About, Contact, FAQ, Shipping and Returns, Privacy, and Terms pages.'),
('Enquiries and notifications','Manage contact-form and product enquiries. Send order confirmations, order status updates and admin new-order alerts through transactional email.'),
('Store settings','Manage store details, logo, contact information, social links, payment configuration and email settings.'),
('Reports and analytics','All reports shall support date-range filtering and Excel or CSV export. Include daily, weekly and monthly sales totals; sales by brand and category; top-selling products; orders by status; payments by method such as M-Pesa versus card and successful versus failed; new and returning customers and top customers; low-stock and out-of-stock products; and coupon usage. Integrate Google Analytics 4 and Meta Facebook Pixel for traffic, visitor behaviour and marketing tracking.')]:e(label,text)
h('5 Customer storefront scope')
for label,text in [
('Header and navigation','Provide menus by brand and category, search with live suggestions, and cart and account icons.'),
('Homepage','Include a hero slider, Kärcher, Midea and LG brand showcase, featured categories, featured or best-selling or new products, promotional banners and newsletter signup.'),
('Brand pages','Provide a dedicated page for each brand with its banner, introduction and complete product listing.'),
('Category and product listings','Filter by brand, category, price and specifications; sort by price, newest and popularity; provide pagination.'),
('Product detail pages','Include image galleries with zoom, regular and offer prices, stock status, key specification tables, descriptions, warranty information, downloadable brochures or manuals, related products and a WhatsApp enquiry button.'),
('Cart and checkout','Allow items to be added, updated or removed; apply coupons and discounts; calculate delivery by location; support guest checkout or login; accept M-Pesa STK Push or card payments and applicable configured methods; provide an order confirmation page and email. Customer payment completion shall automatically update the corresponding transaction and order payment status.'),
('Customer account','Provide registration and login, order history and status tracking, saved addresses, wishlist, profile and password management. Customers shall be able to submit return or refund requests through the frontend.'),
('Information pages','Provide About, Contact with form, map and WhatsApp, FAQ, Shipping and Returns, Privacy Policy, and Terms and Conditions.'),
('General requirements','Support mobile, tablet and desktop layouts, SEO-ready clean URLs and meta tags, fast loading through caching, CDN and optimised images, a WhatsApp chat button and social media links.')]:e(label,text)
doc.save(out)
print(out)
